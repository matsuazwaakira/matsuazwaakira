import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
D = "/home/user/matsuazwaakira/research/data"
OUT = "/home/user/matsuazwaakira/research/ファクトチェック結果_2026-10.xlsx"
rows = json.load(open(f"{D}/factcheck_tohoku.json")) + json.load(open(f"{D}/factcheck_kitakanto.json"))
F = "Arial"; thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
FIX = {("五所川原市", "交通"): "終期をR10に修正（◎→対象外）", ("北秋田市", "環境"): "終期をR8に修正（交通計画で◎は維持）",
       ("福島市", "交通"): "終期をR12に修正（◎→対象外）"}
wb = Workbook(); ws = wb.active; ws.title = "ファクトチェック結果"
ws["A1"] = "終期R9年度と判定した計画のファクトチェック結果（2026年10月1日実施）"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = "当初調査とは別の検索経路で、計画書・計画ページ等の本文に「計画期間」の記述があるかを確認。検索要約1回のみの場合は「確認できず」とした。"
ws["A2"].font = Font(name=F, size=10, color="555555")
H = ["No", "県", "市町村", "計画", "当初の主張", "判定", "確認できた期間", "次期計画の着手状況", "根拠（引用）", "根拠URL", "反映内容", "備考"]
W = [5, 8, 11, 22, 16, 10, 26, 40, 45, 38, 26, 50]
for c, h in enumerate(H, 1):
    x = ws.cell(4, c, h); x.font = Font(name=F, bold=True, color="FFFFFF"); x.fill = PatternFill("solid", fgColor="1F4E78")
    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); ws.column_dimensions[x.column_letter].width = W[c - 1]
color = {"正しい": "C6EFCE", "誤り": "FFC7CE", "確認できず": "FFEB9C"}
for i, r in enumerate(rows):
    rr = 5 + i
    plan_short = "交通" if r["plan"].startswith("交通") else "環境"
    np_ = r.get("next_plan") or ""
    fix = FIX.get((r["name"], plan_short)) or ("「推定」に格下げ" if r["verdict"] == "確認できず" and "：" not in r["plan"] else "")
    if np_.startswith("着手済") and not fix: fix = "改定状況を「着手済」に更新（◎→△）" if r["name"] in ("白石市", "秋田市", "栃木市") else ""
    vals = [i + 1, r["pref"], r["name"], r["plan"], r.get("claimed"), r["verdict"], r.get("found_period"), np_,
            r.get("evidence_quote"), r.get("evidence_url"), fix, r.get("note")]
    for c, v in enumerate(vals, 1):
        x = ws.cell(rr, c, v); x.font = Font(name=F, size=10); x.alignment = Alignment(vertical="top", wrap_text=True); x.border = bd
    v = ws.cell(rr, 6); v.fill = PatternFill("solid", fgColor=color.get(r["verdict"], "FFFFFF")); v.font = Font(name=F, size=10, bold=True)
    v.alignment = Alignment(horizontal="center", vertical="top")
    u = r.get("evidence_url")
    if u and str(u).startswith("http"):
        ws.cell(rr, 10).hyperlink = u; ws.cell(rr, 10).font = Font(name=F, size=9, color="0563C1", underline="single")
last = 4 + len(rows)
ws.freeze_panes = "D5"; ws.auto_filter.ref = f"A4:L{last}"
s = last + 2
ws.cell(s, 2, "集計").font = Font(name=F, bold=True)
for k, lab in enumerate(["正しい", "誤り", "確認できず"]):
    ws.cell(s + 1 + k, 2, lab).font = Font(name=F)
    ws.cell(s + 1 + k, 3, f'=COUNTIF($F$5:$F${last},B{s+1+k})').font = Font(name=F)
ws.cell(s + 4, 2, "合計").font = Font(name=F, bold=True)
ws.cell(s + 4, 3, f"=SUM(C{s+1}:C{s+3})").font = Font(name=F, bold=True)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT); print(OUT, len(rows))
