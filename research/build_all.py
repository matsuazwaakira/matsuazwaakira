import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

S = "/home/user/matsuazwaakira/research/data"
REGION = sys.argv[1] if len(sys.argv) > 1 else "all"
ALL = {
    "青森県": ["aomori1", "aomori2", "aomori3"], "岩手県": ["iwate"], "宮城県": ["miyagi"],
    "秋田県": ["akita"], "山形県": ["yamagata"],
    "福島県": ["fukushima1", "fukushima2", "fukushima3", "fukushima4", "fukushima5"],
    "茨城県": ["ibaraki1", "ibaraki2", "ibaraki3", "ibaraki4"], "栃木県": ["tochigi"],
}
if REGION == "tohoku":
    PREFS = [(p, ALL[p]) for p in ("青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県")]
    OUT = "/home/user/matsuazwaakira/research/東北6県_公共交通計画_環境基本計画_終期一覧_2026-09.xlsx"
    TITLE = "東北6県"
elif REGION == "one":
    PREFS = list(ALL.items())
    OUT = "/home/user/matsuazwaakira/research/自治体営業リスト_統合版_2026-10.xlsx"
    TITLE = "自治体営業リスト（統合版）"
else:
    PREFS = list(ALL.items())
    OUT = "/home/user/matsuazwaakira/research/自治体_公共交通計画_環境基本計画_終期一覧_2026-09.xlsx"
    TITLE = "自治体"

F = "Arial"
thin = Side(style="thin", color="BFBFBF")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
STATUSES = ("未確認", "着手済", "策定済", "計画なし")

wb = Workbook()
cfg = wb.active
cfg.title = "集計・設定"
MIN, MAX = "TGT_MIN", "TGT_MAX"
from openpyxl.workbook.defined_name import DefinedName
for _n, _c in (("TGT_MIN", "$C$4"), ("TGT_MAX", "$C$5")):
    _d = DefinedName(_n, attr_text=f"'集計・設定'!{_c}")
    try:
        wb.defined_names[_n] = _d
    except TypeError:
        wb.defined_names.append(_d)

def load(files):
    rows = []
    for f in files:
        p = f"{S}/{f}.json"
        if os.path.exists(p):
            rows += json.load(open(p))
        else:
            print("MISSING", p)
    return rows

def judge(end, st):
    return (f'=IF({st}="計画なし","－",IF({st}="策定済","✕",IF({end}="","？",'
            f'IF(AND({end}>={MIN},{end}<={MAX}),IF({st}="着手済","△","◎"),"✕"))))')

H = ["No", "市町村", "営業判定", "交通判定", "環境判定",
     "地域公共交通計画", "交通 計画期間", "交通 終期(令和年度)", "交通 確度", "交通 改定状況", "交通 計画リンク（エビデンス）",
     "環境基本計画", "環境 計画期間", "環境 終期(令和年度)", "環境 確度", "環境 改定状況", "環境 計画リンク（エビデンス）",
     "注釈（改定作業の状況など）", "調査状況", "国交省一覧(R8.5末)"]
W = [5, 12, 16, 7, 7, 30, 15, 9, 8, 9, 34, 30, 15, 9, 8, 9, 34, 55, 14, 24]
HR = 3
fills = [("◎", "C6EFCE"), ("△", "FFEB9C"), ("？", "DDEBF7"), ("✕", "F2F2F2")]
ranges = {}

MLIT = json.load(open(f"{S}/mlit_r0805.json"))
_n = lambda x: x.replace("ケ", "ヶ")

def mlit_of(pref, d):
    m = MLIT.get(pref)
    if not m:
        return None
    n = _n(d.get("name") or "")
    own = n in {_n(x) for x in m["own"]}
    wide = [k for k, v in m["wide"].items() if n in {_n(x) for x in v}]
    if not own and not wide:
        lab = "未掲載"
    else:
        lab = "掲載" + ("（広域：" + "・".join(wide) + "）" if wide else "")
        if n in {_n(x) for x in m["expired"]}:
            lab += "／期間満了"
    have = bool(d.get("t_name") or d.get("t_url"))
    if lab != "未掲載" and not have:
        lab += "【当方未把握】"
    if lab == "未掲載" and have:
        lab += "【当方は計画あり】"
    return lab

def status_of(d):
    t = bool(d.get("t_name") or d.get("t_url"))
    e = bool(d.get("e_name") or d.get("e_url") or d.get("e_status") == "計画なし")
    return "調査済" if t and e else "交通のみ調査" if t else "環境のみ調査" if e else "未調査"

def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else None

for pref, files in PREFS:
    data = load(files)
    ws = wb.create_sheet(pref)
    ws["A1"] = f"{pref}　地域公共交通計画・環境基本計画の終期一覧（2026年9月30日時点の調査）"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = ('="営業判定の対象：終期 R"&' + MIN + '&"〜R"&' + MAX + '&"年度 →◎（着手済は△、期間不明は？）。終期セルの色：緑＝対象年度／橙＝その翌年度（次の候補）。対象年度は「集計・設定」シートのC4・C5のプルダウンで切替。"')
    ws["A2"].font = Font(name=F, size=10, color="555555")
    for c, h in enumerate(H, 1):
        cell = ws.cell(HR, c, h)
        cell.font = Font(name=F, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[cell.column_letter].width = W[c - 1]
    ws.row_dimensions[HR].height = 32
    for i, d in enumerate(data):
        r = HR + 1 + i
        ts = d.get("t_status") if d.get("t_status") in STATUSES else "未確認"
        es = d.get("e_status") if d.get("e_status") in STATUSES else "未確認"
        vals = [i + 1, d.get("name"), None, None, None,
                d.get("t_name"), d.get("t_period"), num(d.get("t_end")), d.get("t_conf"), ts, d.get("t_url"),
                d.get("e_name"), d.get("e_period"), num(d.get("e_end")), d.get("e_conf"), es, d.get("e_url"),
                d.get("note"), status_of(d), mlit_of(pref, d)]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(r, c, v)
            cell.font = Font(name=F, size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = bd
        ws.cell(r, 4).value = judge(f"H{r}", f"J{r}")
        ws.cell(r, 5).value = judge(f"N{r}", f"P{r}")
        ws.cell(r, 3).value = (f'=IF(OR(D{r}="◎",E{r}="◎"),"◎ 営業",IF(OR(D{r}="△",E{r}="△"),"△ 着手済・要確認",'
                               f'IF(OR(D{r}="？",E{r}="？"),"？ 期間要確認","✕ 対象外")))')
        ws.cell(r, 3).font = Font(name=F, size=10, bold=True)
        for c in (3, 4, 5, 8, 9, 10, 14, 15, 16):
            ws.cell(r, c).alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        ws.cell(r, 19).alignment = Alignment(horizontal="center", vertical="top")
        if ws.cell(r, 19).value != "調査済":
            ws.cell(r, 19).font = Font(name=F, size=10, bold=True, color="C00000")
        mv = ws.cell(r, 20).value or ""
        if "【" in mv or "期間満了" in mv:
            ws.cell(r, 20).font = Font(name=F, size=10, bold=True, color="C00000")
        for c in (8, 14):
            if ws.cell(r, c + 1).value == "推定":
                ws.cell(r, c).font = Font(name=F, size=10, italic=True, color="7F6000")
        for c in (11, 17):
            u = ws.cell(r, c).value
            if u and str(u).startswith("http"):
                ws.cell(r, c).hyperlink = u
                ws.cell(r, c).font = Font(name=F, size=9, color="0563C1", underline="single")
    last = HR + max(len(data), 1)
    for sym, color in fills:
        ws.conditional_formatting.add(f"A{HR+1}:T{last}",
            FormulaRule(formula=[f'LEFT($C{HR+1},1)="{sym}"'], fill=PatternFill("solid", fgColor=color)))
    for col in ("H", "N"):
        rng = f"{col}{HR+1}:{col}{last}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND(ISNUMBER({col}{HR+1}),{col}{HR+1}>={MIN},{col}{HR+1}<={MAX})'],
            fill=PatternFill("solid", fgColor="63BE7B"), font=Font(bold=True, color="FFFFFF"), stopIfTrue=True))
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'AND(ISNUMBER({col}{HR+1}),{col}{HR+1}={MAX}+1)'],
            fill=PatternFill("solid", fgColor="F4B183"), font=Font(bold=True)))
    _pri = 1
    for _cf in ws.conditional_formatting:
        if str(_cf.sqref).startswith(("H", "N")):
            for _r in _cf.rules:
                _r.priority = _pri; _pri += 1
    for _cf in ws.conditional_formatting:
        if not str(_cf.sqref).startswith(("H", "N")):
            for _r in _cf.rules:
                _r.priority = _pri; _pri += 1
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"J{HR+1}:J{last}")
    dv.add(f"P{HR+1}:P{last}")
    ws.freeze_panes = f"C{HR+1}"
    ws.auto_filter.ref = f"A{HR}:T{last}"
    ranges[pref] = (f"'{pref}'!$C${HR+1}:$C${last}", len(data), f"'{pref}'!$S${HR+1}:$S${last}")

# 集計・設定シート
cfg["A1"] = f"{TITLE} 地域公共交通計画・環境基本計画 終期一覧　集計・設定"
cfg["A1"].font = Font(name=F, bold=True, size=14)
cfg["A3"] = "■ 営業対象とする終期（令和年度）"
cfg["A3"].font = Font(name=F, bold=True)
for rr, lab, v in ((4, "下限", 9), (5, "上限", 9)):
    cfg.cell(rr, 2, lab).font = Font(name=F)
    c = cfg.cell(rr, 3, v)
    c.font = Font(name=F, bold=True, color="0000FF")
    c.fill = PatternFill("solid", fgColor="FFFF00")
yv = DataValidation(type="list", formula1='"7,8,9,10,11,12,13,14,15"', allow_blank=False)
cfg.add_data_validation(yv)
yv.add("C4"); yv.add("C5")
cfg["D4"] = "▼プルダウンで選ぶと全シートの判定（◎）と終期セルの色が切り替わります。例：R10年度に終わる計画も見る→上限を10に／R10だけ見る→下限・上限とも10に。"
cfg["D5"] = "R8年度で終わる計画は作成中のため対象外（初期値は下限9・上限9）。終期セルの色：緑＝対象年度、橙＝上限の翌年度（次の営業候補）。"
cfg["D5"].font = Font(name=F, size=9, color="555555")
cfg["D4"].font = Font(name=F, size=9, color="555555")

cfg["A7"] = "■ 県別の集計"
cfg["A7"].font = Font(name=F, bold=True)
labels = ["◎ 営業", "△ 着手済・要確認", "？ 期間要確認", "✕ 対象外"]
hdr = ["県", "市町村数"] + labels + ["未調査", "一部のみ調査"]
for c, h in enumerate(hdr, 2):
    cell = cfg.cell(8, c, h)
    cell.font = Font(name=F, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center")
for i, (pref, _) in enumerate(PREFS):
    r = 9 + i
    rng, n, srng = ranges[pref]
    cfg.cell(r, 8, f'=COUNTIF({srng},"未調査")').font = Font(name=F)
    cfg.cell(r, 9, f'=COUNTIF({srng},"交通のみ調査")+COUNTIF({srng},"環境のみ調査")').font = Font(name=F)
    cfg.cell(r, 2, pref).font = Font(name=F)
    cfg.cell(r, 3, f'=COUNTA({rng})').font = Font(name=F)
    for k, lab in enumerate(labels):
        cfg.cell(r, 4 + k, f'=COUNTIF({rng},"{lab}")').font = Font(name=F)
tr = 9 + len(PREFS)
cfg.cell(tr, 2, "合計").font = Font(name=F, bold=True)
for c in range(3, 10):
    L = cfg.cell(8, c).column_letter
    cfg.cell(tr, c, f"=SUM({L}9:{L}{tr-1})").font = Font(name=F, bold=True)

notes = [
    ("■ 判定の見方", True),
    ("◎ 営業：どちらかの計画の終期が対象年度（初期値R9）で、次の計画の作業は見つかっていない", False),
    ("△ 着手済・要確認：終期は対象年度だが、次期計画の策定業務委託やパブコメをすでに確認済。委託先が決まっている可能性が高い", False),
    ("？ 期間要確認：終期がわからないため判定できない", False),
    ("✕ 対象外：終期が対象年度の外（R8以前は作成中または終了、R10以降は先）、または新計画を策定済", False),
    ("－：その計画がない（環境基本計画を作っていない町村など）", False),
    ("", False),
    ("■ 入力する列", True),
    ("各県シートの「終期(令和年度)」と「改定状況」（未確認／着手済／策定済／計画なし）を直すと、判定が自動で変わります。", False),
    ("「未確認」は、改定の動きが検索で見つからなかったという意味です。動きがないと確かめたわけではありません。", False),
    ("確度：確認済＝検索結果の本文に期間が書かれていた／推定＝策定時期などから推測（終期のセルは茶色の斜体）／不明＝確認できなかった", False),
    ("", False),
    ("■ 国交省一覧との照合（T列）", True),
    ("国土交通省「地域公共交通計画の作成状況一覧」（令和8年5月末時点）と照合。掲載＝計画作成済み、期間満了＝一覧で灰色（計画期間が満了）。一覧には計画期間の記載がないため、終期は各計画書で確認。", False),
    ("【当方未把握】＝一覧には計画があるが本調査で計画書を特定できていない／【当方は計画あり】＝本調査では計画を把握したが一覧に未掲載（旧法計画・最近の策定・名称違いの可能性）。", False),
    ("", False),
    ("■ 調査の方法と限界", True),
    ("【重要】検索回数に上限（1セッション200回）があるため、県によっては一部の市町村・計画を調べきれていません。「調査状況」列が「未調査」「交通のみ調査」の行は空欄・？になっていますが、計画がないという意味ではありません。", False),
    ("2026年9月30日時点でWeb検索して調べました。自治体サイトやPDFを直接開けない環境だったため、期間は検索結果の本文・要約から取っています。", False),
    ("計画リンクはエビデンスとして検索結果に出たURLを載せています。営業に行く前に、リンク先の計画書で計画期間を確かめてください。特に「推定」「不明」の行は必ず確認してください。", False),
]
for i, (t, b) in enumerate(notes):
    c = cfg.cell(tr + 2 + i, 1, t)
    c.font = Font(name=F, bold=b, size=10)
cfg.column_dimensions["A"].width = 4
for L, w in zip("BCDEFGHI", (12, 10, 12, 18, 16, 12, 10, 14)):
    cfg.column_dimensions[L].width = w

if REGION == "one":
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    from build_memo import fill_memo
    from build_factcheck import fill_fc
    from openpyxl.worksheet.hyperlink import Hyperlink
    ms = wb.create_sheet("営業メモ"); fill_memo(ms)
    fs = wb.create_sheet("ファクトチェック結果"); fill_fc(fs)
    memo_row = {}
    for r in range(5, ms.max_row + 1):
        nm, pf = ms.cell(r, 3).value, ms.cell(r, 2).value
        if not nm or not pf: continue
        for part in str(nm).replace("（", "・").replace("）", "").split("・"):
            memo_row[(pf, part.strip())] = r
        if "避難地域" in str(nm):
            memo_row[("福島県", "*避難*")] = r
    for pref, _ in PREFS:
        ws = wb[pref]
        c = ws.cell(HR, 21, "営業メモ"); c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78"); c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions["U"].width = 10
        for r in range(HR + 1, ws.max_row + 1):
            nm = ws.cell(r, 2).value
            mr = memo_row.get((pref, nm))
            if mr is None and "避難地域" in str(ws.cell(r, 6).value or "") and ("福島県", "*避難*") in memo_row:
                mr = memo_row[("福島県", "*避難*")]
            if mr:
                x = ws.cell(r, 21, "→メモ")
                x.hyperlink = Hyperlink(ref=x.coordinate, location=f"'営業メモ'!C{mr}", display="→メモ")
                x.font = Font(name=F, size=10, color="0563C1", underline="single"); x.alignment = Alignment(horizontal="center", vertical="top")
                x.border = bd
    # 目次
    toc_r = cfg.max_row + 2
    cfg.cell(toc_r, 1, "■ シート一覧（クリックで移動）").font = Font(name=F, bold=True)
    for k, name in enumerate([p for p, _ in PREFS] + ["営業メモ", "ファクトチェック結果"]):
        x = cfg.cell(toc_r + 1 + k, 2, name)
        x.hyperlink = Hyperlink(ref=x.coordinate, location=f"'{name}'!A1", display=name)
        x.font = Font(name=F, size=10, color="0563C1", underline="single")
wb.calculation.fullCalcOnLoad = True
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
print(OUT, {p: ranges[p][1] for p, _ in PREFS})
