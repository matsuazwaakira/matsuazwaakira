import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
D = "/home/user/matsuazwaakira/research/data"
OUT = "/home/user/matsuazwaakira/research/営業メモ_R9終期自治体_2026-10.xlsx"
rows = json.load(open(f"{D}/memos_E.json")) + json.load(open(f"{D}/memos_F.json"))
order = ["青森県", "岩手県", "宮城県", "秋田県", "山形県", "福島県", "茨城県", "栃木県"]
rows.sort(key=lambda r: order.index(r["pref"]) if r["pref"] in order else 99)
F = "Arial"; thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook(); ws = wb.active; ws.title = "営業メモ"
ws["A1"] = "営業メモ：終期R9年度の計画を持つ自治体（2026年10月作成）"; ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = "Web検索で確認できた範囲の情報です。「推定」「要確認」は訪問前に電話・HPで確認してください。次期計画に着手済みの自治体は「着手」列に●。"
ws["A2"].font = Font(name=F, size=10, color="555555")
H = ["No", "県", "自治体", "着手", "対象計画", "計画の概要", "担当課・連絡先", "次期計画の見込みスケジュール", "前回の策定業務（方式・受託者・金額）", "関連する動き", "訪問時の切り口", "根拠URL", "備考"]
W = [4, 7, 12, 5, 24, 46, 28, 46, 38, 40, 42, 40, 30]
HR = 4
for c, h in enumerate(H, 1):
    x = ws.cell(HR, c, h); x.font = Font(name=F, bold=True, color="FFFFFF", size=10); x.fill = PatternFill("solid", fgColor="1F4E78")
    x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); ws.column_dimensions[x.column_letter].width = W[c - 1]
def txt(v):
    if isinstance(v, list): return "\n".join(f"・{s}" for s in v)
    return v
for i, r in enumerate(rows):
    rr = HR + 1 + i
    ns = r.get("next_schedule") or ""
    started = "●" if ("着手済" in ns[:40] or "【次期計画は着手済】" in ns) else ""
    srcs = r.get("sources") or []
    vals = [i + 1, r["pref"], r["name"], started, r.get("target_plan"), r.get("plan_summary"), r.get("dept"), ns,
            r.get("past_contract"), r.get("related"), txt(r.get("talking_points")), "\n".join(srcs), r.get("note")]
    for c, v in enumerate(vals, 1):
        x = ws.cell(rr, c, txt(v)); x.font = Font(name=F, size=9); x.alignment = Alignment(vertical="top", wrap_text=True); x.border = bd
    ws.cell(rr, 3).font = Font(name=F, size=10, bold=True)
    ws.cell(rr, 4).alignment = Alignment(horizontal="center", vertical="top")
    if started:
        ws.cell(rr, 4).font = Font(name=F, size=11, bold=True, color="C00000")
        ws.cell(rr, 8).fill = PatternFill("solid", fgColor="FFF2CC")
    pc = r.get("past_contract") or ""
    if pc and not pc.startswith("不明"):
        ws.cell(rr, 9).fill = PatternFill("solid", fgColor="E2EFDA")
    if srcs and str(srcs[0]).startswith("http"):
        ws.cell(rr, 12).hyperlink = srcs[0]; ws.cell(rr, 12).font = Font(name=F, size=8, color="0563C1", underline="single")
    else:
        ws.cell(rr, 12).font = Font(name=F, size=8)
last = HR + len(rows)
ws.freeze_panes = "D5"; ws.auto_filter.ref = f"A{HR}:M{last}"
ws.print_title_rows = f"{HR}:{HR}"
ws.page_setup.orientation = "landscape"; ws.page_setup.paperSize = ws.PAPERSIZE_A3
ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.cell(last + 2, 2, "凡例").font = Font(name=F, bold=True, size=9)
ws.cell(last + 3, 2, "黄色＝次期計画に着手済み（策定業務の公告等を確認）／緑＝前回の策定業務の情報あり／根拠URL列は先頭URLのみリンク").font = Font(name=F, size=9)
wb.save(OUT); print(OUT, len(rows))
